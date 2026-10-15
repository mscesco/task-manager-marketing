"""A planilha da tela inicial (Spec 058, 15/10) -- o painel em Excel.

Pedido dela: exportar o painel em PDF (isso é da tela, pela impressão) e em
planilha, para quem quer mexer nos números. Quatro abas:

    Resumo     o time, o período, a pessoa, e os números dos cards de cima;
    Gráfico    as colunas (por dia ou por semana): concluídas, no prazo,
               atrasadas;
    Tarefas    as tarefas que o painel CONTOU, com o estado de hoje;
    Carga      a carga por pessoa -- só para quem a vê na tela.

⚠️ OS NÚMEROS SAEM DO MESMO `DashboardService` DA TELA, e as tarefas dos
mesmos predicados (`dashboard_export_rows`): a planilha não pode discordar do
que a pessoa estava olhando quando clicou em exportar.
"""

from __future__ import annotations

import io
import uuid
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy.ext.asyncio import AsyncSession

from app.db.models import Team, User
from app.modules.tasks.application.dashboard_service import DashboardService
from app.modules.tasks.domain.dashboard import TIMEZONE, Period
from app.modules.tasks.infrastructure.task_repository import TaskRepository

_STATE = {"done": "Concluída", "on_time": "No prazo", "overdue": "Atrasada"}
_PRIORITY = {"LOW": "Baixa", "MEDIUM": "Média", "HIGH": "Alta", "URGENT": "Urgente"}
_HEAD = Font(bold=True, color="FFFFFF")
_HEAD_FILL = PatternFill("solid", fgColor="0075DE")
XLSX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _dia(d) -> str:
    return d.strftime("%d/%m/%Y") if d else ""


def _table(ws, header: list[str], rows: list[list]) -> None:
    ws.append(header)
    for cell in ws[1]:
        cell.font = _HEAD
        cell.fill = _HEAD_FILL
    for r in rows:
        ws.append(r)
    for i, titulo in enumerate(header, start=1):
        largura = max([len(str(titulo))] + [len(str(r[i - 1])) for r in rows] + [8])
        ws.column_dimensions[get_column_letter(i)].width = min(largura + 2, 60)
    ws.freeze_panes = "A2"


async def build_dashboard_xlsx(
    session: AsyncSession,
    under_team_id: uuid.UUID | None,
    *,
    assignee_id: uuid.UUID | None = None,
    period: Period | None = None,
) -> tuple[bytes, str]:
    """O arquivo e o nome dele ("painel-marketing-2026-10-12-a-2026-10-18.xlsx")."""
    d = await DashboardService(session).build(
        under_team_id, assignee_id=assignee_id, period=period
    )
    rows = await TaskRepository(session).dashboard_export_rows(
        under_team_id, d.now, assignee_id=assignee_id, period=period
    )
    time_nome = "Todos os times"
    if under_team_id is not None:
        team = await session.get(Team, under_team_id)
        time_nome = team.name if team else time_nome
    pessoa_nome = ""
    if assignee_id is not None:
        user = await session.get(User, assignee_id)
        pessoa_nome = user.name if user else ""
    c = d.counts

    wb = Workbook()
    resumo = wb.active
    resumo.title = "Resumo"
    gerado = datetime.now(TIMEZONE).strftime("%d/%m/%Y %H:%M")
    periodo = (
        f"{_dia(d.period.start)} a {_dia(d.period.end)}"
        if d.period
        else f"Agora (semana de {_dia(c.summary_start)} a {_dia(c.summary_end)})"
    )
    linhas = [
        ["Time", time_nome],
        ["Período", periodo],
        ["Pessoa", pessoa_nome or "Todas"],
        ["Gerado em", gerado],
        ["Estado das tarefas", "como estão hoje"],
        [],
        ["Concluídas", c.done],
        ["Total (concluídas + abertas com prazo)", c.total],
        ["Atrasadas", c.overdue_count],
        ["Em aberto no prazo" if d.period else "Vencem hoje", c.on_time_total],
    ]
    if d.queue is not None:
        linhas.append(["Solicitações pendentes", d.queue.pending])
    for linha in linhas:
        resumo.append(linha)
    for cell in resumo["A"]:
        cell.font = Font(bold=True)
    resumo.column_dimensions["A"].width = 40
    resumo.column_dimensions["B"].width = 40

    grafico = wb.create_sheet("Gráfico")
    _table(
        grafico,
        ["De", "Até", "Concluídas", "No prazo", "Atrasadas"],
        [[_dia(b.start), _dia(b.end), b.done, b.on_time, b.overdue] for b in c.buckets],
    )

    tarefas = wb.create_sheet("Tarefas")
    _table(
        tarefas,
        [
            "Tarefa", "Estado", "Coluna", "Prioridade",
            "Prazo", "Hora", "Responsáveis", "Subtarefa de",
        ],
        [
            [
                r.title,
                _STATE.get(r.state, r.state),
                r.column_name,
                _PRIORITY.get(r.priority, r.priority),
                _dia(r.due_date),
                r.due_time.strftime("%H:%M") if r.due_time else "",
                r.assignees,
                r.parent_title or "",
            ]
            for r in rows
        ],
    )

    if d.workload is not None:
        carga = wb.create_sheet("Carga")
        _table(
            carga,
            ["Pessoa", "Abertas", "Atrasadas", "Conta ativa"],
            [
                [w.name, w.open_count, w.overdue_count, "Sim" if w.is_active else "Não"]
                for w in d.workload
            ],
        )

    buf = io.BytesIO()
    wb.save(buf)
    janela = (
        f"{d.period.start.isoformat()}-a-{d.period.end.isoformat()}"
        if d.period
        else d.now.today.isoformat()
    )
    nome = f"painel-{_slug(time_nome)}-{janela}.xlsx"
    return buf.getvalue(), nome


def _slug(s: str) -> str:
    import unicodedata

    base = unicodedata.normalize("NFKD", s).encode("ascii", "ignore").decode()
    return "-".join("".join(ch if ch.isalnum() else " " for ch in base.lower()).split()) or "painel"
