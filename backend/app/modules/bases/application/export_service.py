"""Exportar uma Base em planilha (Spec 057, 15/10).

⚠️ A TELA MANDA AS LINHAS E AS COLUNAS, NA ORDEM DA VISAO: filtro e ordenacao
da visao moram no navegador (`lib/baseViews.ts`), e reescrever aqui seria uma
segunda regra que divergiria da primeira. O servidor so confere que cada id e
desta base -- id de fora e ignorado, e nao vira erro.

O texto de cada celula e o que a tela mostra (rotulo da opcao, nome da pessoa,
data `14/10/2026`), e os separadores sao os que a IMPORTACAO entende (`; `):
exportar e importar de volta da o mesmo.
"""

from __future__ import annotations

import csv
import io
import re
import uuid
from datetime import date
from typing import Any

from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.tenant import require_tenant
from app.db.models import User
from app.modules.bases.application.access import visible_base
from app.modules.bases.domain.columns import live_options
from app.modules.bases.domain.spreadsheet import FORMULA_START
from app.modules.bases.infrastructure.base_repository import (
    BaseColumnRepository,
    BaseRowRepository,
    BaseTableRepository,
)
from app.shared.exceptions.base import ValidationError

XLSX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
CSV_MEDIA_TYPE = "text/csv; charset=utf-8"


class BaseExportService:
    def __init__(self, session: AsyncSession) -> None:
        self._session = session
        self._bases = BaseTableRepository(session)
        self._columns = BaseColumnRepository(session)
        self._rows = BaseRowRepository(session)

    async def export(
        self,
        base_id: uuid.UUID,
        *,
        fmt: str,
        row_ids: list[uuid.UUID] | None = None,
        column_ids: list[uuid.UUID] | None = None,
    ) -> tuple[bytes, str, str]:
        """O arquivo, o tipo e o nome. Quem le a base, exporta (`base.read`)."""
        if fmt not in ("xlsx", "csv"):
            raise ValidationError("Formato: xlsx ou csv.", details={"field": "format"})
        base = await visible_base(self._bases, base_id)
        colunas = await self._columns.list_for(base.id)
        # ⚠️ `is not None`, e nao "tem algo": a visao com filtro que nao casa
        # nada manda `[]`, e lista vazia lida como "sem filtro" exportava a
        # base inteira -- inclusive o que a visao esconde.
        if column_ids is not None:
            por_id = {c.id: c for c in colunas}
            colunas = [por_id[i] for i in column_ids if i in por_id]
        linhas = await self._rows.list_for(base.id)
        if row_ids is not None:
            por_id_l = {r.id: r for r in linhas}
            linhas = [por_id_l[i] for i in row_ids if i in por_id_l]
        nomes = await self._nomes(linhas, colunas)

        cabecalho = [_texto(c.name, fmt) for c in colunas]
        corpo = [
            [_texto(_celula(c, r.values.get(str(c.id)), nomes, fmt), fmt) for c in colunas]
            for r in linhas
        ]
        nome = _slug(base.name)
        if fmt == "csv":
            buf = io.StringIO()
            # ⚠️ `;` e BOM: e o que o Excel em portugues abre sem perguntar nada.
            w = csv.writer(buf, delimiter=";")
            w.writerow(cabecalho)
            w.writerows(corpo)
            return ("﻿" + buf.getvalue()).encode("utf-8"), CSV_MEDIA_TYPE, f"{nome}.csv"

        wb = Workbook()
        ws = wb.active
        ws.title = _aba(base.name)
        ws.append(cabecalho)
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="0075DE")
        for linha in corpo:
            ws.append(linha)
        # O openpyxl le texto que comeca com `=` como FORMULA: `=HYPERLINK(...)`
        # digitado numa celula virava link vivo no arquivo do gestor, e
        # `=== Secao ===` virava arquivo "corrompido". Texto fica texto.
        for row in ws.iter_rows():
            for cell in row:
                if cell.data_type == "f":
                    cell.data_type = "s"
        for i, c in enumerate(colunas, start=1):
            largura = max([len(c.name)] + [len(str(r[i - 1] or "")) for r in corpo[:200]] + [8])
            ws.column_dimensions[get_column_letter(i)].width = min(largura + 2, 60)
            if c.type == "date":
                for (cell,) in ws.iter_rows(min_row=2, min_col=i, max_col=i):
                    cell.number_format = "DD/MM/YYYY"
        ws.freeze_panes = "A2"
        out = io.BytesIO()
        wb.save(out)
        return out.getvalue(), XLSX_MEDIA_TYPE, f"{nome}.xlsx"

    async def _nomes(self, linhas, colunas) -> dict[str, str]:
        ids: set[str] = set()
        for c in colunas:
            if c.type == "person":
                for r in linhas:
                    v = r.values.get(str(c.id))
                    if isinstance(v, list):
                        ids.update(str(x) for x in v)
        if not ids:
            return {}
        tenant = require_tenant()
        rows = (
            await self._session.execute(
                select(User.id, User.name).where(
                    User.workspace_id == tenant.workspace_id,
                    User.id.in_([uuid.UUID(i) for i in ids]),
                )
            )
        ).all()
        return {str(u): n for u, n in rows}


def _celula(coluna: Any, valor: Any, nomes: dict[str, str], fmt: str) -> Any:
    if valor is None:
        return None
    tipo = coluna.type
    if tipo == "date" and isinstance(valor, str):
        d = date.fromisoformat(valor)
        # No Excel, data DE VERDADE (filtra e ordena); no CSV, como a tela.
        return d if fmt == "xlsx" else d.strftime("%d/%m/%Y")
    if tipo == "checkbox":
        return "Sim" if valor else "Não"
    if tipo == "number":
        return valor
    if tipo in ("select", "multi_select"):
        rotulos = {o["id"]: o["label"] for o in live_options(list(coluna.options))}
        ids = valor if isinstance(valor, list) else [valor]
        return "; ".join(rotulos[i] for i in ids if i in rotulos) or None
    if tipo == "person":
        ids = valor if isinstance(valor, list) else [valor]
        return "; ".join(nomes.get(str(i), "Pessoa removida") for i in ids) or None
    if tipo == "link":
        itens = [{"url": valor}] if isinstance(valor, str) else valor
        return "; ".join(i.get("url", "") for i in itens if isinstance(i, dict)) or None
    return str(valor)


def _texto(valor: Any, fmt: str) -> Any:
    """Texto que o Excel aceita e nao executa."""
    if not isinstance(valor, str):
        return valor
    # Caractere de controle (o `\x0b` que vem colado do PowerPoint) o openpyxl
    # recusa com erro -- e o export inteiro dava 500.
    valor = ILLEGAL_CHARACTERS_RE.sub("", valor)
    # No CSV nao ha como marcar "isto e texto": o apostrofo e a marca do
    # proprio Excel, e a importacao o tira na volta (`spreadsheet.to_text`).
    if fmt == "csv" and valor[:1] in FORMULA_START:
        return "'" + valor
    return valor


def _aba(nome: str) -> str:
    """Nome de aba valido: ate 31 caracteres, sem `/ \\ ? * : [ ]` e sem
    apostrofo nas pontas -- "Pautas 2026/1" dava 500 no export."""
    limpo = re.sub(r"[\\/?*:\[\]]", " ", nome)[:31].strip().strip("'").strip()
    return limpo or "Base"


def _slug(s: str) -> str:
    import unicodedata

    base = unicodedata.normalize("NFKD", s).encode("ascii", "ignore").decode()
    return "-".join("".join(ch if ch.isalnum() else " " for ch in base.lower()).split()) or "base"
