"""Planilha <-> Base (Spec 057, 15/10) -- ler o arquivo e converter valores.

Puro, sem banco: o que e uma data, um numero ou uma caixinha numa planilha
brasileira, e qual tipo de coluna uma coluna do arquivo PARECE ser. Quem sabe
de opcoes e de pessoas (que dependem do banco) e o `import_service`.

⚠️ AS DUAS FORMAS DE ARQUIVO CHEGAM DIFERENTES: o Excel ja traz tipo (data e
`datetime`, numero e `float`); o CSV traz SO texto. Por isso todo conversor
aceita os dois -- e texto no formato brasileiro (`14/10/2026`, `1.234,5`).
"""

from __future__ import annotations

import csv
import io
import math
import re
from collections.abc import Iterator
from dataclasses import dataclass
from datetime import date, datetime
from typing import Any

from openpyxl import load_workbook

from app.shared.exceptions.base import ValidationError

#: O arquivo maximo (o corpo chega em base64, ~1,37x maior).
MAX_FILE_BYTES = 5 * 1024 * 1024
#: Quantas linhas a previa mostra.
PREVIEW_ROWS = 20
#: Colunas lidas do arquivo -- o mesmo teto do mapeamento (`_ImportColumn.index`).
MAX_COLUMNS = 500
#: Linhas lidas do arquivo CONTANDO AS EM BRANCO. ⚠️ O Excel guarda linha so
#: formatada como linha: sem este teto, uma planilha "vazia" de 1 milhao de
#: linhas prende o processo por minutos antes de qualquer conferencia.
MAX_RAW_ROWS = 50_000
#: Separador de varios valores numa celula (selecao multipla, pessoa).
#: ⚠️ A VIRGULA TAMBEM (15/10): o Google Planilhas copia e exporta os chips de
#: selecao multipla como "Formulario, Import, Site" -- sem ela, tudo virava UMA
#: opcao com o texto inteiro.
MULTI_SEP = re.compile(r"\s*[;,\n]\s*")

#: Primeiro caractere que faz o Excel ler a celula como formula.
FORMULA_START = "=+-@"

_DATA_BR = re.compile(r"^(\d{1,2})/(\d{1,2})/(\d{2}|\d{4})$")
_DATA_ISO = re.compile(r"^(\d{4})-(\d{2})-(\d{2})")
_MILHAR = re.compile(r"^-?\d{1,3}(\.\d{3})+$")
_SIM = {"sim", "s", "x", "true", "verdadeiro", "1", "yes", "y", "ok", "✓", "✔"}
_NAO = {"nao", "não", "n", "false", "falso", "0", "no", ""}


@dataclass(frozen=True)
class Table:
    sheets: list[str]
    sheet: str
    header: list[str]
    rows: list[list[Any]]


def read_table(
    content: bytes,
    filename: str,
    *,
    sheet: str | None = None,
    header_row: int = 1,
    max_rows: int | None = None,
) -> Table:
    """O arquivo como cabecalho + linhas. `header_row` e 1-based (a linha do
    cabecalho; as de cima sao ignoradas -- planilha com titulo em cima).

    `max_rows`: mais linhas com conteudo do que isso e recusa ANTES de ler o
    resto do arquivo. ⚠️ E CPU pura (o openpyxl leva segundos numa planilha
    grande): quem chama de rota async roda isto em `asyncio.to_thread`."""
    if len(content) > MAX_FILE_BYTES:
        raise ValidationError("O arquivo passa de 5 MB.", details={"field": "file"})
    if header_row < 1:
        raise ValidationError("A linha do cabecalho comeca em 1.", details={"field": "header_row"})
    nome = filename.lower()
    if nome.endswith(".xlsx") or nome.endswith(".xlsm"):
        sheets, escolhida, linhas = _ler_xlsx(content, sheet)
    elif nome.endswith(".csv") or nome.endswith(".txt"):
        sheets, escolhida, linhas = ["CSV"], "CSV", _ler_csv(content)
    else:
        raise ValidationError(
            "Envie uma planilha .xlsx ou .csv.", details={"field": "file"}
        )
    try:
        cabecalho, corpo = _corpo(linhas, header_row, max_rows)
    except csv.Error as exc:  # aspas sem fechar, campo gigante
        raise ValidationError(
            "Nao consegui ler o CSV (aspas sem fechar?).", details={"field": "file"}
        ) from exc
    finally:
        close = getattr(linhas, "close", None)
        if close:
            close()  # fecha a planilha mesmo quando a leitura para no meio
    return Table(sheets=sheets, sheet=escolhida, header=cabecalho, rows=corpo)


def _corpo(
    linhas: Iterator[list[Any]], header_row: int, max_rows: int | None
) -> tuple[list[str], list[list[Any]]]:
    for _ in range(header_row - 1):
        next(linhas, None)
    primeira = next(linhas, None)
    if primeira is None:
        raise ValidationError("A planilha esta vazia.", details={"field": "file"})
    cabecalho = [str(c).strip() if c is not None else "" for c in primeira]
    largura = len(cabecalho)
    corpo: list[list[Any]] = []
    for lidas, linha in enumerate(linhas, start=1):
        if lidas > MAX_RAW_ROWS:
            raise ValidationError(
                f"A planilha tem mais de {_milhar(MAX_RAW_ROWS)} linhas, contando as em "
                "branco. Apague as linhas vazias do fim e tente de novo.",
                details={"field": "file"},
            )
        linha = list(linha) + [None] * (largura - len(linha))
        linha = linha[:largura]
        if all(_vazio(v) for v in linha):
            continue  # linha em branco no meio do arquivo nao vira linha na base
        if max_rows is not None and len(corpo) >= max_rows:
            raise ValidationError(
                f"A planilha passa de {_milhar(max_rows)} linhas.",
                details={"field": "file", "limit": max_rows},
            )
        corpo.append(linha)
    # Coluna sem nome no cabecalho ganha um ("Coluna 3"), em vez de sumir.
    cabecalho = [h or f"Coluna {i + 1}" for i, h in enumerate(cabecalho)]
    return cabecalho, corpo


def _ler_xlsx(
    content: bytes, sheet: str | None
) -> tuple[list[str], str, Iterator[list[Any]]]:
    try:
        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception as exc:  # arquivo corrompido ou que nao e xlsx
        raise ValidationError(
            "Nao consegui abrir a planilha.", details={"field": "file"}
        ) from exc
    nomes = list(wb.sheetnames)
    escolhida = sheet if sheet in nomes else nomes[0]
    ws = wb[escolhida]
    # ⚠️ `max_col` COMPLETA a linha ate ele: com o teto puro, toda planilha
    # viria com 500 colunas. A largura de verdade, limitada pelo teto.
    largura = min(ws.max_column or MAX_COLUMNS, MAX_COLUMNS)

    def linhas() -> Iterator[list[Any]]:
        try:
            for r in ws.iter_rows(values_only=True, max_col=largura):
                yield list(r)
        finally:
            wb.close()

    return nomes, escolhida, linhas()


def _ler_csv(content: bytes) -> Iterator[list[Any]]:
    # ⚠️ utf-8 (com ou sem BOM) primeiro. Depois cp1252, que e o que o Excel em
    # portugues grava: travessao, aspas curvas e "€" sao letras nele e
    # caracteres de controle no latin-1 -- e o latin-1 nunca falha, entao o
    # texto entrava estragado sem aviso. latin-1 fica de ultima rede (o cp1252
    # recusa 5 bytes que o latin-1 aceita).
    for codificacao in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            texto = content.decode(codificacao)
            break
        except UnicodeDecodeError:
            continue
    amostra = texto[:4096]
    # O Excel em pt-BR separa por `;`; o resto do mundo, por `,`.
    sep = ";" if amostra.count(";") >= amostra.count(",") else ","
    return (r[:MAX_COLUMNS] for r in csv.reader(io.StringIO(texto), delimiter=sep))


def _milhar(n: int) -> str:
    return f"{n:,}".replace(",", ".")


def _vazio(v: Any) -> bool:
    return v is None or (isinstance(v, str) and not v.strip())


# ------------------------------------------------------------ conversao


class ConversionError(Exception):
    """Um valor que nao serve para o tipo -- vira erro DA LINHA na previa."""


def to_text(v: Any) -> str | None:
    if _vazio(v):
        return None
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    if isinstance(v, datetime | date):
        return to_date(v)  # data numa coluna de texto sai como a tela mostra
    s = str(v).strip()
    # O apostrofo que o EXPORT poe na frente de `= + - @` (para o Excel nao ler
    # como formula) sai na volta: exportar e importar da o mesmo.
    if len(s) > 1 and s[0] == "'" and s[1] in FORMULA_START:
        return s[1:]
    return s


def to_number(v: Any) -> int | float | None:
    if _vazio(v):
        return None
    if isinstance(v, bool):
        raise ConversionError("Numero esperado.")
    if isinstance(v, int | float):
        if isinstance(v, float) and not math.isfinite(v):
            raise ConversionError("Numero invalido.")
        return int(v) if isinstance(v, float) and v.is_integer() else v
    s = str(v).strip().replace("R$", "").replace("%", "").replace(" ", "")
    # "1.234,5" (pt-BR) e "1234.5" (ponto decimal): a virgula decide.
    # ⚠️ E "1.200" SEM virgula e MIL E DUZENTOS, e nao 1,2 (o teste pegou):
    # ponto seguido de grupos de 3 digitos e separador de milhar.
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    elif _MILHAR.match(s):
        s = s.replace(".", "")
    try:
        n = float(s)
    except ValueError as exc:
        raise ConversionError(f"'{v}' nao e um numero.") from exc
    return int(n) if n.is_integer() else n


def to_date(v: Any) -> str | None:
    """`YYYY-MM-DD`, que e como a celula de data guarda."""
    if _vazio(v):
        return None
    if isinstance(v, datetime):
        return v.date().isoformat()
    if isinstance(v, date):
        return v.isoformat()
    s = str(v).strip()
    m = _DATA_BR.match(s)
    if m:
        dia, mes, ano = int(m[1]), int(m[2]), int(m[3])
        ano = ano + 2000 if ano < 100 else ano
        try:
            return date(ano, mes, dia).isoformat()
        except ValueError as exc:
            raise ConversionError(f"'{s}' nao e uma data.") from exc
    m = _DATA_ISO.match(s)
    if m:
        try:
            return date(int(m[1]), int(m[2]), int(m[3])).isoformat()
        except ValueError as exc:
            raise ConversionError(f"'{s}' nao e uma data.") from exc
    raise ConversionError(f"'{s}' nao e uma data (use 14/10/2026).")


def to_checkbox(v: Any) -> bool | None:
    if v is None:
        return None
    if isinstance(v, bool):
        return v
    if isinstance(v, int | float):
        return v != 0
    s = str(v).strip().casefold()
    if s in _SIM:
        return True
    if s in _NAO:
        return False
    raise ConversionError(f"'{v}' nao e sim/nao.")


def split_many(v: Any) -> list[str]:
    """Varios valores numa celula: separados por `;`, `,` ou quebra de linha."""
    if _vazio(v):
        return []
    return [p for p in MULTI_SEP.split(str(v).strip()) if p]


def to_links(v: Any) -> list[dict[str, str]] | None:
    # ⚠️ Link NAO separa por virgula: ela aparece dentro de URL.
    if _vazio(v):
        return None
    partes = [p for p in re.split(r"[;\s]+", str(v).strip()) if p]
    return [{"url": p} for p in partes] or None


# ------------------------------------------------------------ sugestao de tipo


def suggest_type(values: list[Any]) -> str:
    """O tipo que a coluna do arquivo PARECE ter -- so uma sugestao: a pessoa
    escolhe no assistente."""
    cheios = [v for v in values if not _vazio(v)]
    if not cheios:
        return "text"

    def todos(conv) -> bool:
        try:
            for v in cheios:
                conv(v)
        except ConversionError:
            return False
        return True

    if all(isinstance(v, bool) for v in cheios):
        return "checkbox"
    if todos(to_date) and not all(isinstance(v, int | float) for v in cheios):
        return "date"
    if todos(to_number):
        return "number"
    textos = [str(v).strip() for v in cheios]
    if all(t.casefold() in _SIM | _NAO for t in textos):
        return "checkbox"
    if all(re.match(r"^https?://", t) for t in textos):
        return "link"
    if any(MULTI_SEP.search(t) for t in textos):
        # Só parece seleção múltipla se as PARTES são curtas e se repetem
        # (chips), e não frases com vírgula.
        partes = [p for t in textos for p in split_many(t)]
        distintas = {p.casefold() for p in partes}
        if len(distintas) <= 30 and sum(map(len, partes)) / max(1, len(partes)) <= 25:
            return "multi_select"
        return "text"
    distintos = {t.casefold() for t in textos}
    if len(distintos) <= 20 and len(distintos) <= max(2, len(textos) // 2):
        return "select"
    return "text"
